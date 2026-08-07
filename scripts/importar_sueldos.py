import openpyxl
import sqlite3
from collections import defaultdict
import datetime
import os

print("=== INICIANDO SCRIPT DE IMPORTACIÓN Y AUDITORÍA DE SUELDOS CON CENTRO & DICCIONARIO REFACTORIZADO ===")

excel_path = r'SUELDO202605_DefAjustado_escuelas_minimizado.xlsx'
excel_ref_path = r'datos_csv/CENTROS Y SECTORES EDUCACION REFACTORIZADO (3).XLSX'
db_path = r'database/database.sqlite'

if not os.path.exists(excel_path):
    print(f"ERROR: No se encuentra el archivo Excel en {excel_path}")
    exit(1)

# 0. Cargar Diccionario Maestro Refactorizado (CENTRO + SECTOR)
dict_refactorizado = {}
if os.path.exists(excel_ref_path):
    print(f"Cargando Diccionario Maestro Refactorizado desde {excel_ref_path}...")
    wb_ref = openpyxl.load_workbook(excel_ref_path, data_only=True, read_only=True)
    ws_ref = wb_ref.active
    for row in ws_ref.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or row[3] is None:
            continue
        try:
            c = int(float(row[0]))
            s = int(float(row[3]))
            dict_refactorizado[(c, s)] = {
                'nivel': str(row[1] or '').strip(),
                'gestion': str(row[2] or '').strip(),
                'nom_sector': str(row[4] or '').strip(),
                'nom_centro': str(row[5] or '').strip(),
            }
        except Exception:
            pass
    wb_ref.close()
    print(f"Diccionario Maestro cargado con {len(dict_refactorizado)} mapeos de (CENTRO, SECTOR).")
else:
    print("Aviso: No se encontró el archivo refactorizado. Se procederá sin diccionario secundario.")

# 1. Asegurar columna 'centro' en la base de datos SQLite
conn = sqlite3.connect(db_path)
cursor = conn.cursor()

cursor.execute("PRAGMA table_info(auditoria_radio_resultados)")
cols = [row[1] for row in cursor.fetchall()]
if 'centro' not in cols:
    print("Añadiendo columna 'centro' a auditoria_radio_resultados...")
    cursor.execute("ALTER TABLE auditoria_radio_resultados ADD COLUMN centro INTEGER")
    conn.commit()

# 2. Cargar modalidades de SIGE desde database.sqlite
print("Cargando modalidades SIGE desde la base de datos...")
sige_query = """
SELECT 
    m.id as modalidad_id,
    CAST(m.sector AS INTEGER) as sector_int,
    m.sector as sector_str,
    m.radio as radio_sige,
    m.zona as zona_sige,
    m.inst_legal_radio,
    m.nivel_educativo,
    m.ambito,
    e.cue,
    e.nombre as nombre_establecimiento,
    ed.cui,
    ed.localidad,
    ed.radio_circ,
    ed.radio_camino
FROM modalidades m
JOIN establecimientos e ON e.id = m.establecimiento_id
JOIN edificios ed ON ed.id = e.edificio_id
WHERE m.deleted_at IS NULL AND e.deleted_at IS NULL AND ed.deleted_at IS NULL
"""
sige_rows = cursor.execute(sige_query).fetchall()

# Indexar modalidades SIGE por sector_int
modalidades_by_sector = defaultdict(list)
for r in sige_rows:
    mod = {
        'modalidad_id': r[0],
        'sector': r[1],
        'radio_sige': r[3],
        'zona_sige': r[4],
        'inst_legal_radio': r[5],
        'nivel_educativo': r[6],
        'ambito': r[7],
        'cue': r[8],
        'nombre_establecimiento': r[9],
        'cui': r[10],
        'localidad': r[11],
        'radio_circ': r[12],
        'radio_camino': r[13],
    }
    if r[1] is not None and r[1] > 0:
        modalidades_by_sector[r[1]].append(mod)

print(f"Modalidades SIGE cargadas: {len(sige_rows)} ({len(modalidades_by_sector)} sectores distintos en SIGE)")

# 3. Leer Excel de Sueldos
print(f"Abriendo {excel_path}...")
wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
ws = wb.active

IDX_CENTRO = 2
IDX_SECTOR = 3
IDX_CLASE  = 4
IDX_ZONA   = 20
IDX_A01    = 21
IDX_A04    = 23

grupos = defaultdict(lambda: {
    'filas': 0,
    'a01_list': [],
    'a04_list': [],
    'clases': set(),
})

total_filas_excel = 0
print("Procesando filas de sueldos...")

for row in ws.iter_rows(min_row=2, values_only=True):
    total_filas_excel += 1
    c = row[IDX_CENTRO]
    s = row[IDX_SECTOR]
    z = row[IDX_ZONA]

    if c is None or s is None:
        continue

    try:
        centro = int(c)
        sector = int(s)
    except ValueError:
        continue

    zona = str(z).strip() if z is not None else 'S/D'

    key = (centro, sector, zona)
    grupos[key]['filas'] += 1

    clase = row[IDX_CLASE]
    if clase is not None:
        grupos[key]['clases'].add(int(clase))

    a01 = row[IDX_A01]
    a04 = row[IDX_A04]

    if a01 is not None and a04 is not None:
        try:
            v_a01 = float(a01)
            v_a04 = float(a04)
            if v_a01 > 0 and v_a04 >= 0:
                grupos[key]['a01_list'].append(v_a01)
                grupos[key]['a04_list'].append(v_a04)
        except ValueError:
            pass

wb.close()
print(f"Total filas procesadas del Excel: {total_filas_excel}")
print(f"Grupos unívocos (CENTRO, SECTOR, ZONA) creados: {len(grupos)}")

# 4. Calcular mediana y determinar radio para cada grupo
def calcular_mediana(lista):
    if not lista:
        return None
    s = sorted(lista)
    n = len(s)
    if n % 2 == 1:
        return s[n // 2]
    else:
        return (s[n // 2 - 1] + s[n // 2]) / 2.0

def calcular_radio_sueldo(porcentaje):
    if porcentaje is None:
        return None
    if porcentaje <= 45.0:
        return 1
    elif porcentaje <= 55.0:
        return 2
    elif porcentaje <= 85.0:
        return 3
    elif porcentaje <= 105.0:
        return 4
    elif porcentaje <= 125.0:
        return 5
    elif porcentaje <= 145.0:
        return 6
    else:
        return 7

# Asegurar o buscar nomina_id
cursor.execute("SELECT id FROM nominas_sueldos WHERE periodo = '2026-05' LIMIT 1")
row_nom = cursor.fetchone()
if row_nom:
    nomina_id = row_nom[0]
else:
    now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    cursor.execute("""
        INSERT INTO nominas_sueldos (periodo, archivo_nombre, total_filas, total_sectores, fecha_importacion, created_at, updated_at)
        VALUES ('2026-05', 'SUELDO202605_DefAjustado_escuelas_minimizado.xlsx', ?, ?, ?, ?, ?)
    """, (total_filas_excel, len(grupos), now_str, now_str, now_str))
    nomina_id = cursor.lastrowid
    conn.commit()

# Limpiar resultados anteriores para esta nomina
cursor.execute("DELETE FROM auditoria_radio_resultados WHERE nomina_id = ?", (nomina_id,))
conn.commit()

# 5. Mapear y generar auditoría
print("Generando registros de auditoría...")
records_to_insert = []
now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

CENTROS_PRIVADOS = {98}

for (centro, sector, zona_sueldo), datos in grupos.items():
    filas_docentes = datos['filas']
    
    porc_mediana = None
    if datos['a01_list'] and datos['a04_list']:
        porcentajes_filas = [
            (a4 / a1) * 100.0 for a1, a4 in zip(datos['a01_list'], datos['a04_list']) if a1 > 0
        ]
        porc_mediana = calcular_mediana(porcentajes_filas)

    if porc_mediana is not None:
        porc_mediana = round(porc_mediana, 2)
        radio_sueldo = calcular_radio_sueldo(porc_mediana)
    else:
        radio_sueldo = None

    # Buscar coincidencia en SIGE (Prioridad 1)
    mods_sige = modalidades_by_sector.get(sector, [])
    
    selected_mod = None
    if mods_sige:
        if len(mods_sige) == 1:
            selected_mod = mods_sige[0]
        else:
            es_privado = (centro in CENTROS_PRIVADOS)
            candidates = []
            for m in mods_sige:
                m_priv = 'privad' in str(m['ambito']).lower()
                if es_privado == m_priv:
                    candidates.append(m)
            if candidates:
                selected_mod = candidates[0]
            else:
                selected_mod = mods_sige[0]

    notas_auditor = None

    if selected_mod:
        cue = selected_mod['cue']
        nombre_est = selected_mod['nombre_establecimiento']
        nivel_edu = selected_mod['nivel_educativo']
        localidad = selected_mod['localidad']
        zona_sige = selected_mod['zona_sige']
        radio_sige = selected_mod['radio_sige']
        radio_circ = selected_mod['radio_circ']
        radio_camino = selected_mod['radio_camino']
        inst_legal_radio = selected_mod['inst_legal_radio']
    else:
        cue = None
        nombre_est = None
        nivel_edu = None
        localidad = None
        zona_sige = None
        radio_sige = None
        radio_circ = None
        radio_camino = None
        inst_legal_radio = None

        # Prioridad 2: Enriquecer con Diccionario Refactorizado del Usuario si no hay SIGE
        ref = dict_refactorizado.get((centro, sector))
        if ref:
            nombre_est = ref['nom_sector'] if ref['nom_sector'] else None
            nivel_edu = ref['nivel'] if ref['nivel'] else 'GENERAL'
            if ref['nom_centro']:
                notas_auditor = f"Identificado: {ref['nom_centro']} ({ref['gestion']})"

    coincide_zona = 1 if (zona_sige and zona_sueldo and zona_sige.upper() == zona_sueldo.upper()) else 0

    # Clasificar estado_auditoria
    if not selected_mod:
        estado_auditoria = 'SIN_SIGE'
    elif radio_sueldo is None or radio_sige is None:
        estado_auditoria = 'SIN_RADIO_PAGADO'
    elif radio_sueldo == radio_sige and (radio_circ is not None and radio_sueldo == radio_circ) and (radio_camino is not None and radio_sueldo == radio_camino):
        estado_auditoria = 'COINCIDE_TOTAL'
    elif radio_sueldo == radio_sige and (radio_camino is not None and radio_sueldo == radio_camino):
        estado_auditoria = 'COINCIDE_SIGE_Y_CAMINO'
    elif radio_sueldo == radio_sige and (radio_circ is not None and radio_sueldo == radio_circ):
        estado_auditoria = 'COINCIDE_SIGE_Y_CIRC'
    elif radio_sueldo == radio_sige:
        estado_auditoria = 'COINCIDE_SIGE'
    elif radio_sueldo > radio_sige:
        estado_auditoria = 'PAGA_MAS_QUE_SIGE'
    else:
        estado_auditoria = 'PAGA_MENOS_QUE_SIGE'

    records_to_insert.append((
        nomina_id,
        centro,
        sector,
        zona_sueldo,
        zona_sige,
        coincide_zona,
        nivel_edu,
        nombre_est,
        cue,
        localidad,
        radio_sueldo,
        radio_sige,
        radio_circ,
        radio_camino,
        0, # radio_justificado
        inst_legal_radio,
        porc_mediana,
        'NUEVA',
        filas_docentes,
        estado_auditoria,
        'PENDIENTE',
        notas_auditor,
        now_str,
        now_str
    ))

print(f"Insertando {len(records_to_insert)} registros en auditoria_radio_resultados...")
cursor.executemany("""
    INSERT INTO auditoria_radio_resultados (
        nomina_id, centro, sector, zona_sueldo, zona_sige, coincide_zona,
        nivel_educativo, nombre_establecimiento, cue, localidad,
        radio_sueldo, radio_sige, radio_circ, radio_camino, radio_justificado,
        inst_legal_radio, porc_pagado_mediana, escala_usada, total_filas_docentes,
        estado_auditoria, estado_gestion, notas_auditor, created_at, updated_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
""", records_to_insert)

cursor.execute("""
    UPDATE nominas_sueldos
    SET total_filas = ?, total_sectores = ?, updated_at = ?
    WHERE id = ?
""", (total_filas_excel, len(records_to_insert), now_str, nomina_id))

conn.commit()

print("\n=== RESUMEN DE AUDITORÍA CON DICCIONARIO REFACTORIZADO DE CENTROS Y SECTORES ===")
cursor.execute("""
    SELECT estado_auditoria, COUNT(*) as cant
    FROM auditoria_radio_resultados
    WHERE nomina_id = ?
    GROUP BY estado_auditoria
    ORDER BY cant DESC
""", (nomina_id,))
for row in cursor.fetchall():
    print(f"  {row[0]}: {row[1]}")

conn.close()
print("\n=== PROCESO COMPLETADO EXITOSAMENTE Y SIN RUIDO ===")

# Ejecutar auditoría y depuración automática de centros y sectores
try:
    import auditar_centros_sectores
    auditar_centros_sectores.ejecutar_auditoria()
except Exception as e:
    print(f"Aviso: No se pudo ejecutar la auditoría de depuración automática ({e})")

