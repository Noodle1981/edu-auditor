import sqlite3
import openpyxl
import os
import json

def audit_sprint1_kpi():
    print("======================================================================")
    print("=== SPRINT 1: AUDITORÍA QA DE 5 NIVELES - PESTAÑA RESUMEN & KPIS ===")
    print("======================================================================")

    f_sueldos = r'SUELDO202605_DefAjustado_escuelas_minimizado.xlsx'
    f_master = r'datos_csv/CENTROS Y SECTORES EDUCACION REFACTORIZADO (3).XLSX'
    db_app = r'database/database.sqlite'
    db_sige = r'database/establecimientos.sqlite'

    results = {}

    # ------------------------------------------------------------------
    # NIVEL 1: ARCHIVO EXCEL DE SUELDOS (FUENTE PRIMARIA DE HABERES)
    # ------------------------------------------------------------------
    print("\n[NIVEL 1] Leyendo Fuente de Haberes Excel 'SUELDO202605_DefAjustado_escuelas_minimizado.xlsx'...")
    wb_s = openpyxl.load_workbook(f_sueldos, data_only=True, read_only=True)
    ws_s = wb_s['ESCU']
    rows_s = ws_s.iter_rows(values_only=True)
    header_s = next(rows_s)

    idx_centro = header_s.index('CENTRO')
    idx_sector = header_s.index('SECTOR')
    idx_a04 = header_s.index('A04') if 'A04' in header_s else -1

    excel_filas_totales = 0
    excel_centros_set = set()
    excel_sectores_set = set()
    excel_pairs_count = {}
    excel_centros_filas = {}

    for r in rows_s:
        excel_filas_totales += 1
        c, s = r[idx_centro], r[idx_sector]
        if c is not None and s is not None:
            try:
                c_int = int(c)
                s_int = int(s)
                excel_centros_set.add(c_int)
                excel_sectores_set.add(s_int)
                pair = (c_int, s_int)
                excel_pairs_count[pair] = excel_pairs_count.get(pair, 0) + 1
                excel_centros_filas[c_int] = excel_centros_filas.get(c_int, 0) + 1
            except:
                pass
    wb_s.close()

    print(f"  -> Total filas de haberes docentes en Excel: {excel_filas_totales}")
    print(f"  -> Total Centros únicos en Excel: {len(excel_centros_set)}")
    print(f"  -> Total (Centro, Sector) combinaciones únicas con pagos en Excel: {len(excel_pairs_count)}")

    # ------------------------------------------------------------------
    # NIVEL 2: ARCHIVO EXCEL MAESTRO REFACTORIZADO
    # ------------------------------------------------------------------
    print("\n[NIVEL 2] Leyendo Diccionario Maestro Excel 'CENTROS Y SECTORES EDUCACION REFACTORIZADO (3).XLSX'...")
    wb_m = openpyxl.load_workbook(f_master, data_only=True, read_only=True)
    ws_m = wb_m['establecimientos']
    rows_m = ws_m.iter_rows(values_only=True)
    header_m = next(rows_m)

    m_idx_c = header_m.index('centro')
    m_idx_s = header_m.index('SECTOR')

    master_centros_set = set()
    master_pairs_set = set()

    for r in rows_m:
        c, s = r[m_idx_c], r[m_idx_s]
        if c is not None and s is not None:
            try:
                c_int = int(c)
                s_int = int(s)
                master_centros_set.add(c_int)
                master_pairs_set.add((c_int, s_int))
            except:
                pass
    wb_m.close()

    print(f"  -> Total combinaciones en catálogo Maestro: {len(master_pairs_set)} ({len(master_centros_set)} centros)")

    # ------------------------------------------------------------------
    # NIVEL 3: BASE DE DATOS DE REFERENCIA SIGE (establecimientos.sqlite)
    # ------------------------------------------------------------------
    print("\n[NIVEL 3] Leyendo DB Referencia SIGE 'database/establecimientos.sqlite'...")
    conn_sige = sqlite3.connect(db_sige)
    cursor_sige = conn_sige.cursor()

    cant_est_sige = cursor_sige.execute("SELECT COUNT(*) FROM establecimientos").fetchone()[0]
    cant_mod_sige = cursor_sige.execute("SELECT COUNT(*) FROM modalidades WHERE deleted_at IS NULL").fetchone()[0]
    cant_sectores_sige = cursor_sige.execute("SELECT COUNT(DISTINCT CAST(sector AS INTEGER)) FROM modalidades WHERE deleted_at IS NULL AND sector IS NOT NULL AND sector > 0").fetchone()[0]
    conn_sige.close()

    print(f"  -> Establecimientos registrados en SIGE: {cant_est_sige}")
    print(f"  -> Modalidades activas en SIGE: {cant_mod_sige}")
    print(f"  -> Sectores numéricos únicos en SIGE: {cant_sectores_sige}")

    # ------------------------------------------------------------------
    # NIVEL 4: BASE DE DATOS DE APLICACIÓN (database/database.sqlite)
    # ------------------------------------------------------------------
    print("\n[NIVEL 4] Auditando DB Aplicación 'database/database.sqlite'...")
    conn_app = sqlite3.connect(db_app)
    cursor_app = conn_app.cursor()

    # A. Tabla auditoria_radio_resultados
    db_arr_rows = cursor_app.execute("SELECT COUNT(*), SUM(total_filas_docentes), COUNT(DISTINCT centro), COUNT(DISTINCT sector) FROM auditoria_radio_resultados").fetchone()
    arr_cant_registros = db_arr_rows[0]
    arr_sum_docentes = db_arr_rows[1]
    arr_centros = db_arr_rows[2]
    arr_sectores = db_arr_rows[3]

    print(f"  -> Tabla 'auditoria_radio_resultados': {arr_cant_registros} sectores auditados | Suma docentes: {arr_sum_docentes} | Centros: {arr_centros} | Sectores: {arr_sectores}")

    # B. Sectores vinculados a CUE vs desvinculados
    linked_count = cursor_app.execute("SELECT COUNT(*), SUM(total_filas_docentes) FROM auditoria_radio_resultados WHERE cue IS NOT NULL AND cue != ''").fetchone()
    unlinked_count = cursor_app.execute("SELECT COUNT(*), SUM(total_filas_docentes) FROM auditoria_radio_resultados WHERE cue IS NULL OR cue = ''").fetchone()

    print(f"     * Vinculados a CUE (Escuelas): {linked_count[0]} sectores ({linked_count[1]} docentes)")
    print(f"     * Desvinculados (Sin CUE): {unlinked_count[0]} sectores ({unlinked_count[1]} docentes)")

    # C. Tabla auditoria_sueldo_registros_viejos
    viejos_stats = cursor_app.execute("SELECT COUNT(*), COUNT(DISTINCT sector) FROM auditoria_sueldo_registros_viejos").fetchone()
    print(f"  -> Tabla 'auditoria_sueldo_registros_viejos': {viejos_stats[0]} registros de escala vieja ({viejos_stats[1]} sectores)")

    # D. Tabla depuracion_centros_sectores
    depuracion_stats = cursor_app.execute("SELECT estado_depuracion, COUNT(*) FROM depuracion_centros_sectores GROUP BY estado_depuracion").fetchall()
    print(f"  -> Tabla 'depuracion_centros_sectores': {dict(depuracion_stats)}")

    # ------------------------------------------------------------------
    # NIVEL 5: CÁLCULOS DEL BACKEND LARAVEL (AuditoriaSueldosController)
    # ------------------------------------------------------------------
    print("\n[NIVEL 5] Auditando Lógica de Backend (AuditoriaSueldosController.php) y KPIs...")

    # Réplica exacta de la consulta de AuditoriaSueldosController index()
    query_resultados = """
    SELECT 
        r.id, r.centro, r.sector, r.total_filas_docentes, r.estado_auditoria, r.cue,
        COALESCE(e.nombre, r.nombre_establecimiento) as nombre_establecimiento
    FROM auditoria_radio_resultados r
    LEFT JOIN establecimientos e ON e.cue = r.cue
    GROUP BY r.id
    ORDER BY r.sector
    """
    resultados_backend = cursor_app.execute(query_resultados).fetchall()

    # Conteo de linkedResultados (solo los que tienen CUE)
    linked_resultados = [r for r in resultados_backend if r[5] is not None and str(r[5]).strip() != '']
    
    total_centros_kpi = len(set(r[1] for r in resultados_backend if r[1] is not None))
    total_sectores_kpi = len(set(r[2] for r in linked_resultados if r[2] is not None))
    total_docentes_kpi = sum(r[3] or 0 for r in linked_resultados)

    coinciden_sige_kpi = sum(r[3] or 0 for r in linked_resultados if r[4] in ('COINCIDE_TOTAL', 'COINCIDE_SIGE', 'COINCIDE_SIGE_Y_CAMINO', 'COINCIDE_SIGE_Y_CIRC'))
    tasa_coincidencia = round((coinciden_sige_kpi / total_docentes_kpi * 100), 1) if total_docentes_kpi > 0 else 0

    paga_mas_sectores = len([r for r in linked_resultados if r[4] == 'PAGA_MAS_QUE_SIGE'])
    paga_mas_docentes = sum(r[3] or 0 for r in linked_resultados if r[4] == 'PAGA_MAS_QUE_SIGE')

    paga_menos_sectores = len([r for r in linked_resultados if r[4] == 'PAGA_MENOS_QUE_SIGE'])
    paga_menos_docentes = sum(r[3] or 0 for r in linked_resultados if r[4] == 'PAGA_MENOS_QUE_SIGE')

    sin_sige_sectores = len([r for r in linked_resultados if r[4] == 'SIN_SIGE'])
    sin_sige_docentes = sum(r[3] or 0 for r in linked_resultados if r[4] == 'SIN_SIGE')

    print("  -> KPIs Calculados por Controller para la Web UI:")
    print(f"     • Total Centros: {total_centros_kpi} (Coincide con Excel: {total_centros_kpi == len(excel_centros_set)})")
    print(f"     • Total Sectores (Linked): {total_sectores_kpi}")
    print(f"     • Total Filas Docentes (Linked): {total_docentes_kpi}")
    print(f"     • % Coincidencia de Radio: {tasa_coincidencia}% ({coinciden_sige_kpi} de {total_docentes_kpi} docentes)")
    print(f"     • Pagan MÁS que SIGE: {paga_mas_sectores} sectores ({paga_mas_docentes} docentes)")
    print(f"     • Pagan MENOS que SIGE: {paga_menos_sectores} sectores ({paga_menos_docentes} docentes)")
    print(f"     • Sin Registro SIGE: {sin_sige_sectores} sectores ({sin_sige_docentes} docentes)")

    # ------------------------------------------------------------------
    # MATRIZ DE CONSISTENCIA Y AUDITORÍA DE INCONSISTENCIAS DE DATOS
    # ------------------------------------------------------------------
    print("\n======================================================================")
    print("=== MATRIZ DE EVALUACIÓN DE INCONSISTENCIAS (SPRINT 1) ===")
    print("======================================================================")

    inconsistencias = []

    # Check 1: ¿La suma total de filas docentes en DB coincide con la suma del Excel?
    if arr_sum_docentes == excel_filas_totales:
        print(" [OK] Check 1: Las filas totales cargadas en DB (67.321) coinciden 100% con el Excel origen.")
    else:
        inc = f"Desfasaje de filas totales: Excel={excel_filas_totales} vs DB={arr_sum_docentes}"
        print(f" [X] INCONSISTENCIA 1: {inc}")
        inconsistencias.append(inc)

    # Check 2: ¿Los centros únicos en Excel coinciden con la DB y con el KPI?
    if len(excel_centros_set) == arr_centros == total_centros_kpi:
        print(f" [OK] Check 2: Los 22 Centros únicos coinciden al 100% en Excel, SQLite y la vista KPI.")
    else:
        inc = f"Discrepancia en Centros: Excel={len(excel_centros_set)}, DB={arr_centros}, KPI={total_centros_kpi}"
        print(f" [X] INCONSISTENCIA 2: {inc}")
        inconsistencias.append(inc)

    # Check 3: Suma de docentes vinculados + desvinculados debe dar la suma total
    if (linked_count[1] + unlinked_count[1]) == arr_sum_docentes:
        print(f" [OK] Check 3: La suma de docentes vinculados a escuelas ({linked_count[1]}) + desvinculados ({unlinked_count[1]}) da exactamente el total de haberes ({arr_sum_docentes}).")
    else:
        inc = f"Suma de docentes desfasada: Vinculados={linked_count[1]} + Desvinculados={unlinked_count[1]} != Total={arr_sum_docentes}"
        print(f" [X] INCONSISTENCIA 3: {inc}")
        inconsistencias.append(inc)

    # Check 4: Desglose por Centro en KPI vs DB
    print("\nAuditando desglose por Centro en el KPI (`centrosBreakdown`):")
    cursor_app.execute("""
        SELECT centro, COUNT(*), SUM(total_filas_docentes)
        FROM auditoria_radio_resultados
        GROUP BY centro
        ORDER BY SUM(total_filas_docentes) DESC
    """)
    db_centros_breakdown = cursor_app.fetchall()
    
    suma_breakdown_docentes = 0
    for c_row in db_centros_breakdown:
        c_id, c_sectores, c_docentes = c_row
        excel_doc = excel_centros_filas.get(c_id, 0)
        coincide = (c_docentes == excel_doc)
        suma_breakdown_docentes += c_docentes
        status_str = "OK" if coincide else "DESFASADO"
        print(f"   - Centro {c_id:2d}: DB={c_docentes:5d} docentes | Excel={excel_doc:5d} docentes -> [{status_str}]")
        if not coincide:
            inconsistencias.append(f"Centro {c_id} desfasado: DB={c_docentes} vs Excel={excel_doc}")

    conn_app.close()

    print("\n----------------------------------------------------------------------")
    if not inconsistencias:
        print("=== RESULTADO SPRINT 1: PESTAÑA RESUMEN & KPIS APROBADA SIN ERRORES ===")
    else:
        print(f"=== ATENCIÓN: SE DETECTARON {len(inconsistencias)} INCONSISTENCIAS EN SPRINT 1 ===")
    print("----------------------------------------------------------------------")

if __name__ == '__main__':
    audit_sprint1_kpi()
